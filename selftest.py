#!/usr/bin/env python3
"""
Headless self-test for DRT Studio.

Exercises the whole pipeline without needing a display:
  * synthetic spectrum with known circuit parameters -> DRT recovers them
  * all three inversion modes
  * automatic lambda selection (L-curve, GCV)
  * Kramers-Kronig test on clean and drifting data
  * file import: CSV / TSV / Excel, separate Re and Im files, decimal comma
  * peak detection and assignment
  * project save / load round trip
  * PDF, Word and Excel export for a single sample and for a comparison
  * every GUI module is importable and its callbacks resolve

Run:  python selftest.py
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import traceback

import matplotlib
matplotlib.use("Agg")
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from drt_studio.core import io_data
from drt_studio.core.drt import compute_drt
from drt_studio.core.kk import lin_kk
from drt_studio.core.model import MODES, DRTSettings, Project, Sample
from drt_studio.core.peaks import analyse, compare_results, interpret, ringing_score
from drt_studio.core.pipeline import (analyse_project, analyse_sample, build_demo_project,
                                      export_comparison_bundle, export_sample_bundle,
                                      render_comparison_figures, render_sample_figures)
from drt_studio.core.synth import make_spectrum

PASS, FAIL = [], []


def check(name, cond, detail=""):
    if cond:
        PASS.append(name)
        print("  PASS  %-58s %s" % (name, detail))
    else:
        FAIL.append(name)
        print("  FAIL  %-58s %s" % (name, detail))
    return cond


def section(t):
    print("\n" + t)
    print("-" * 78)


def main():
    tmp = tempfile.mkdtemp(prefix="drt_selftest_")
    try:
        return _run(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _run(tmp):
    # ------------------------------------------------------------------ 1
    section("1. DRT inversion against a known circuit")
    true = dict(r_inf=6.0, film_r=12.0, cat_r=8.0, an_r=12.0, diff_r=180.0)
    r_pol_true = true["film_r"] + true["cat_r"] + true["an_r"] + true["diff_r"]
    f, zr, zi = make_spectrum(noise_pct=0.3, seed=1)
    check("synthetic spectrum generated", f.size == 61, "%d points" % f.size)

    st = DRTSettings()
    res = {}
    for m in MODES:
        r = compute_drt(f, zr, zi, st, m)
        analyse(r)
        res[m] = r
        print("        %-9s R_inf=%7.2f  R_pol=%8.2f  RMS=%5.2f%%  peaks=%d"
              % (m, r.r_inf, r.r_pol, r.rms_pct, len(r.peaks)))

    c = res["combined"]
    check("combined: R_inf recovered", abs(c.r_inf - true["r_inf"]) < 2.0,
          "%.2f vs %.1f" % (c.r_inf, true["r_inf"]))
    check("combined: R_pol recovered", abs(c.r_pol - r_pol_true) / r_pol_true < 0.10,
          "%.1f vs %.0f" % (c.r_pol, r_pol_true))
    check("combined: fit is good", c.rms_pct < 2.0, "RMS %.2f %%" % c.rms_pct)
    check("all modes produce a distribution",
          all(res[m].gamma.size > 50 and res[m].gamma.max() > 0 for m in MODES))
    check("gamma is non-negative", all(res[m].gamma.min() >= -1e-9 for m in MODES))
    check("dominant peak near the diffusion tau",
          any(0.1 < p.tau < 2.0 for p in c.peaks),
          "peaks at " + ", ".join("%.3g" % p.tau for p in c.peaks))
    dom = max(c.peaks, key=lambda p: p.resistance)
    check("dominant peak is the diffusion process", dom.label.startswith("P6"),
          "%s, R=%.1f ohm" % (dom.label, dom.resistance))
    check("peak areas sum to about R_pol",
          abs(sum(p.resistance for p in c.peaks) - c.r_pol) / c.r_pol < 0.30,
          "%.1f vs %.1f" % (sum(p.resistance for p in c.peaks), c.r_pol))

    # ------------------------------------------------------------------ 2
    section("2. Automatic regularisation selection")
    for mode in ("lcurve", "gcv"):
        s2 = DRTSettings()
        s2.lambda_mode = mode
        s2.lambda_steps = 12
        r = compute_drt(f, zr, zi, s2, "combined")
        analyse(r)
        check("lambda by %s" % mode, r.lambda_used > 0 and r.rms_pct < 5,
              "lambda=%.3g RMS=%.2f%%" % (r.lambda_used, r.rms_pct))

    section("2b. Regularisation extremes behave as documented")
    s_lo = DRTSettings(); s_lo.lambda_value = 1e-9
    s_hi = DRTSettings(); s_hi.lambda_value = 1e2
    r_lo = analyse(compute_drt(f, zr, zi, s_lo, "re"))
    r_hi = analyse(compute_drt(f, zr, zi, s_hi, "re"))
    check("under-regularised gives more peaks than over-regularised",
          len(r_lo.peaks) >= len(r_hi.peaks),
          "%d vs %d peaks" % (len(r_lo.peaks), len(r_hi.peaks)))
    sc, _ = ringing_score(r_lo)
    check("ringing detector returns a score", 0.0 <= sc <= 1.0, "score %.2f" % sc)

    # ------------------------------------------------------------------ 3
    section("3. Kramers-Kronig validation")
    kk_clean = lin_kk(f, zr, zi)
    check("clean spectrum passes KK", kk_clean.max_abs_pct < 1.0,
          "max %.3f %%" % kk_clean.max_abs_pct)
    fd, zrd, zid = make_spectrum(noise_pct=0.5, drift_pct=25.0, seed=4)
    kk_drift = lin_kk(fd, zrd, zid)
    check("drifting spectrum is flagged", kk_drift.max_abs_pct > kk_clean.max_abs_pct,
          "max %.3f %% vs %.3f %%" % (kk_drift.max_abs_pct, kk_clean.max_abs_pct))

    # ------------------------------------------------------------------ 4
    section("4. File import")
    d = tmp
    # 4a full CSV
    p_full = os.path.join(d, "cellX_full.csv")
    with open(p_full, "w") as fh:
        fh.write("Frequency (Hz),Z' (ohm),Z'' (ohm)\n")
        for a, b, cc in zip(f, zr, zi):
            fh.write("%.8g,%.8g,%.8g\n" % (a, b, cc))
    df = io_data.read_table(p_full)
    cm = io_data.guess_columns(df, "full")
    got = io_data.extract(df, cm)
    check("CSV with headers read", got["f"].size == f.size and cm.freq == 0 and cm.re == 1,
          "%d rows, cols %s/%s/%s" % (len(df), cm.freq, cm.re, cm.im))

    # 4b separate Re / Im, tab separated, -Z'' convention
    p_re = os.path.join(d, "cellY_re.txt")
    p_im = os.path.join(d, "cellY_im.txt")
    with open(p_re, "w") as fh:
        fh.write("freq\tZreal\n")
        for a, b in zip(f, zr):
            fh.write("%.8g\t%.8g\n" % (a, b))
    with open(p_im, "w") as fh:
        fh.write("freq\t-Zimag\n")
        for a, b in zip(f, zi):
            fh.write("%.8g\t%.8g\n" % (a, -b))
    d_re = io_data.extract(io_data.read_table(p_re), io_data.guess_columns(
        io_data.read_table(p_re), "re"))
    dfi = io_data.read_table(p_im)
    cmi = io_data.guess_columns(dfi, "im")
    d_im = io_data.extract(dfi, cmi)
    check("negated imaginary column detected", cmi.im_is_negated,
          "column '%s'" % list(dfi.columns)[cmi.im])
    fm, zrm, zim, how = io_data.merge_re_im(d_re["f"], d_re["z_re"], d_im["f"], d_im["z_im"])
    check("separate Re/Im files merge", fm.size == f.size, how)
    check("merged Im matches original", np.allclose(zim, zi, rtol=1e-6, atol=1e-9))

    # 4c mismatched frequency axes
    f2 = f[::2] * 1.0001
    zi2 = np.interp(np.log10(f2), np.log10(f), zi)
    fm2, zrm2, zim2, how2 = io_data.merge_re_im(f, zr, f2, zi2)
    check("mismatched axes interpolate", fm2.size > 10, how2)

    # 4d decimal comma, semicolon separated
    p_eu = os.path.join(d, "cellZ_eu.csv")
    with open(p_eu, "w") as fh:
        fh.write("f;Zre;Zim\n")
        for a, b, cc in zip(f[:20], zr[:20], zi[:20]):
            fh.write(("%.6g;%.6g;%.6g\n" % (a, b, cc)).replace(".", ","))
    dfe = io_data.read_table(p_eu)
    check("decimal comma handled", len(dfe) == 20 and dfe.shape[1] == 3,
          "shape %s" % (dfe.shape,))

    # 4e Excel
    try:
        import pandas as pd
        p_xl = os.path.join(d, "cellW.xlsx")
        pd.DataFrame({"f/Hz": f, "Z1": zr, "Z2": zi}).to_excel(p_xl, index=False)
        dfx = io_data.read_table(p_xl)
        cmx = io_data.guess_columns(dfx, "full")
        check("Excel file read", len(dfx) == f.size and cmx.freq == 0,
              "sheets %s" % io_data.list_sheets(p_xl))
    except Exception as ex:
        check("Excel file read", False, str(ex))

    # 4f headerless numeric file
    p_nh = os.path.join(d, "nohdr.dat")
    np.savetxt(p_nh, np.c_[f, zr, zi])
    dfn = io_data.read_table(p_nh)
    check("headerless file read", dfn.shape == (f.size, 3), "shape %s" % (dfn.shape,))

    # 4g DRT curve import
    p_drt = os.path.join(d, "curve_drt.csv")
    with open(p_drt, "w") as fh:
        fh.write("tau,gamma\n")
        for t, g in zip(c.tau, c.gamma):
            fh.write("%.8g,%.8g\n" % (t, g))
    dfd = io_data.read_table(p_drt)
    cmd = io_data.guess_columns(dfd, "drt")
    check("DRT curve import", cmd.tau == 0 and cmd.gamma == 1)

    # 4h import-dialog group builder
    from drt_studio.gui.import_dialog import _build_sample, _strip_role_suffix
    check("filename role suffix stripped", _strip_role_suffix("cellY_re") == "cellY",
          _strip_role_suffix("cellY_re"))
    rows = [
        dict(path=p_re, role="re", group="cellY", sheet="", sheets=[],
             df=io_data.read_table(p_re),
             cm=io_data.guess_columns(io_data.read_table(p_re), "re"), neg=False,
             funit="Hz", error=""),
        dict(path=p_im, role="im", group="cellY", sheet="", sheets=[], df=dfi,
             cm=cmi, neg=True, funit="Hz", error=""),
    ]
    s_merged = _build_sample("cellY", rows)
    check("import wizard merges a Re+Im group", s_merged.has_eis and s_merged.n_points > 10,
          "%d points" % s_merged.n_points)
    check("merged sample has correct Im sign", float(np.mean(s_merged.z_im)) < 0,
          "mean Z'' = %.2f" % float(np.mean(s_merged.z_im)))

    # frequency unit conversion
    rows_khz = [dict(path=p_full, role="full", group="k", sheet="", sheets=[], df=df,
                     cm=cm, neg=False, funit="kHz", error="")]
    s_khz = _build_sample("k", rows_khz)
    check("frequency unit conversion", abs(s_khz.f.max() - f.max() * 1e3) < 1e-3,
          "%.4g Hz" % s_khz.f.max())

    # ------------------------------------------------------------------ 5
    section("5. Project, analysis and persistence")
    proj = Project(name="selftest")
    s1 = proj.add(Sample(name="cell 1", f=f, z_re=zr, z_im=zi))
    f2b, zr2, zi2b = make_spectrum(noise_pct=0.4, diff_r=90.0, seed=5)
    s2b = proj.add(Sample(name="cell 2", f=f2b, z_re=zr2, z_im=zi2b))
    analyse_project(proj)
    check("project analysed", all(len(s.results) == 3 for s in proj.samples))
    check("KK ran for every sample", all(s.kk is not None for s in proj.samples))
    check("duplicate names are made unique",
          proj.add(Sample(name="cell 1")).name == "cell 1 (2)")
    proj.remove("cell 1 (2)")

    txt = interpret(s1.results["combined"])
    check("interpretation produced", len(txt) >= 4, "%d findings" % len(txt))
    check("cross-check produced", len(compare_results(s1.results)) >= 2)

    pth = os.path.join(d, "p.drtproj")
    proj.save(pth)
    p2 = Project.load(pth)
    check("project round trip", len(p2.samples) == 2 and
          abs(p2.samples[0].results["combined"].r_pol -
              s1.results["combined"].r_pol) < 1e-6,
          "%.3f ohm" % p2.samples[0].results["combined"].r_pol)
    check("peaks survive round trip",
          len(p2.samples[0].results["combined"].peaks) == len(s1.results["combined"].peaks))

    # ------------------------------------------------------------------ 6
    section("6. Figures")
    figs = render_sample_figures(s1, os.path.join(d, "figs"))
    check("per-sample figures rendered", len(figs) >= 7,
          ", ".join(sorted(figs)[:8]))
    check("figure files are non-trivial",
          all(os.path.getsize(p) > 8000 for p in figs.values()))
    cfigs = render_comparison_figures(proj.samples, os.path.join(d, "cfigs"))
    check("comparison figures rendered", len(cfigs) >= 6, ", ".join(sorted(cfigs)))

    # ------------------------------------------------------------------ 7
    section("7. Export bundles")
    out1 = os.path.join(d, "exp1")
    made = export_sample_bundle(s1, proj, out1, want_pdf=True, want_docx=True,
                                want_excel=True, want_figures=True)
    pdfs = [m for m in made if m.endswith(".pdf")]
    docs = [m for m in made if m.endswith(".docx")]
    xls = [m for m in made if m.endswith(".xlsx")]
    # reportlab / python-docx / openpyxl are optional extras.  Each format is
    # verified when its package is available and reported as unavailable
    # otherwise - the bundle must not fail either way.
    def _opt(label, mod, files, minsize):
        try:
            __import__(mod)
        except ImportError:
            check("%s skipped (no %s)" % (label, mod), not files,
                  "optional dependency absent; nothing written, as expected")
            return
        check("%s written" % label,
              bool(files) and os.path.getsize(files[0]) > minsize,
              "%d bytes" % (os.path.getsize(files[0]) if files else 0))

    _opt("sample PDF", "reportlab", pdfs, 40000)
    _opt("sample Word", "docx", docs, 20000)
    _opt("sample Excel", "openpyxl", xls, 15000)

    out2 = os.path.join(d, "exp2")
    made2 = export_comparison_bundle(proj, proj.samples, out2, want_pdf=True,
                                     want_docx=True, want_excel=True)
    for label, mod, ext in (("comparison PDF", "reportlab", ".pdf"),
                            ("comparison Word", "docx", ".docx"),
                            ("comparison Excel", "openpyxl", ".xlsx")):
        got = any(m.endswith(ext) for m in made2)
        try:
            __import__(mod)
        except ImportError:
            check("%s skipped (no %s)" % (label, mod), not got,
                  "optional dependency absent")
            continue
        check("%s written" % label, got)

    # workbook contents
    try:
        from openpyxl import load_workbook
        if not xls:
            raise ImportError("openpyxl produced no workbook")
        wb = load_workbook(xls[0])
        need = ["Overview"]
        check("workbook has the expected sheets",
              all(n in wb.sheetnames for n in need) and len(wb.sheetnames) >= 4,
              ", ".join(wb.sheetnames[:6]))
        ov = wb["Overview"]
        check("workbook overview is populated", ov.max_row > 4 and ov.max_column >= 6,
              "%dx%d" % (ov.max_row, ov.max_column))
    except Exception as ex:
        check("workbook inspected", False, str(ex))

    # PDF page count.  reportlab is optional, so a missing PDF is reported as a
    # skip rather than crashing the suite (the bundle writes SKIPPED.txt then).
    if not pdfs:
        check("PDF report skipped cleanly",
              any(os.path.basename(m) == "SKIPPED.txt" for m in made),
              "reportlab not installed; bundle recorded the skip and still "
              "produced its other outputs")
    else:
        try:
            import pypdfium2 as pdfium
            doc = pdfium.PdfDocument(pdfs[0])
            check("PDF has a full report", len(doc) >= 6, "%d pages" % len(doc))
        except ImportError:
            with open(pdfs[0], "rb") as fh:
                head = fh.read(5)
            check("PDF is a valid file", head == b"%PDF-", "header %s" % head)

    # ------------------------------------------------------------------ 8
    section("8. Demo project")
    demo = build_demo_project()
    check("demo project built", len(demo.samples) == 4,
          ", ".join(s.name.split(" - ")[0] for s in demo.samples))
    analyse_project(demo)
    check("demo analysed", all("combined" in s.results for s in demo.samples))
    drifting = [s for s in demo.samples if "drift" in s.name.lower()]
    if drifting:
        others = [s for s in demo.samples if s not in drifting]
        check("demo drifting cell has the worst KK residual",
              drifting[0].kk.max_abs_pct > max(o.kk.max_abs_pct for o in others),
              "%.2f %% vs %.2f %%" % (drifting[0].kk.max_abs_pct,
                                      max(o.kk.max_abs_pct for o in others)))
    thin = [s for s in demo.samples if "thin" in s.name.lower()]
    base = [s for s in demo.samples if "baseline" in s.name.lower()]
    if thin and base:
        check("demo thin-cathode cell has lower diffusion resistance",
              thin[0].results["combined"].region_r.get("P6", 0) <
              base[0].results["combined"].region_r.get("P6", 0),
              "%.1f vs %.1f ohm" % (thin[0].results["combined"].region_r.get("P6", 0),
                                    base[0].results["combined"].region_r.get("P6", 0)))

    # ------------------------------------------------------------------ 9
    section("9. GUI modules")
    try:
        import tkinter  # noqa
        have_tk = True
    except Exception:
        have_tk = False
    try:
        import drt_studio.gui.widgets as W
        import drt_studio.gui.import_dialog as ID
        import drt_studio.gui.app as A
        check("GUI modules import", True, "tkinter available: %s" % have_tk)
        for name in ("import_data", "run_current", "run_all", "do_compare",
                     "export_current", "export_all", "export_comparison", "load_demo",
                     "run_kk_current", "run_lcurve", "save_project", "open_project"):
            if not hasattr(A.App, name):
                check("App.%s exists" % name, False)
                break
        else:
            check("all App callbacks defined", True, "12 checked")
        check("PlotPane/SortableTree/ScrollFrame present",
              all(hasattr(W, n) for n in ("PlotPane", "SortableTree", "ScrollFrame",
                                          "LabeledEntry", "Tooltip")))
        check("ImportDialog present", hasattr(ID, "ImportDialog"))
    except Exception as ex:
        check("GUI modules import", False, str(ex))
        traceback.print_exc()

    # ------------------------------------------------------------------ 10
    section("10. Regression: window areas, screen fit, DRT-only samples")

    # (a) window resistances must conserve area (masking used to lose ~4 %)
    try:
        from drt_studio.core.peaks import region_resistances
        r = base[0].results["combined"]
        tot = sum(region_resistances(r).values())
        err = abs(tot - r.r_pol) / max(r.r_pol, 1e-12) * 100
        check("window areas sum to R_pol", err < 0.5,
              "sum=%.3f vs R_pol=%.3f (%.3f %% off)" % (tot, r.r_pol, err))
    except Exception as ex:
        check("window areas sum to R_pol", False, str(ex))

    # (b) fit_to_screen must never exceed the display
    try:
        from drt_studio.gui.widgets import fit_to_screen
        check("fit_to_screen exists", callable(fit_to_screen))
    except Exception as ex:
        check("fit_to_screen exists", False, str(ex))

    # (c) no hard-coded oversized geometry left in the GUI
    try:
        import re as _re2
        bad = []
        for mod in ("app.py", "import_dialog.py"):
            src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "drt_studio", "gui", mod)).read()
            for m in _re2.finditer(r'geometry\("(\d+)x(\d+)"\)', src):
                if int(m.group(1)) > 1000 or int(m.group(2)) > 700:
                    bad.append("%s: %sx%s" % (mod, m.group(1), m.group(2)))
        check("no oversized hard-coded geometry", not bad, ", ".join(bad) or "all clamped")
    except Exception as ex:
        check("no oversized hard-coded geometry", False, str(ex))

    # (d) an imported DRT curve must be usable, not a dead end
    try:
        from drt_studio.gui.import_dialog import _build_sample
        import pandas as _pd
        tf = os.path.join(tmp, "curve.csv")
        _tau = np.logspace(-5, 1, 240)
        _g = 40 * np.exp(-0.5 * ((np.log10(_tau) + 0.5) / 0.25) ** 2)
        _pd.DataFrame({"tau": _tau, "gamma": _g}).to_csv(tf, index=False)
        df = io_data.read_table(tf)
        row = dict(path=tf, role="drt", group="curve", sheet=None, sheets=[], df=df,
                   cm=io_data.guess_columns(df, "drt"), neg="auto", funit="Hz", error="")
        s_drt = _build_sample("curve", [row])
        ok = (s_drt is not None and "combined" in s_drt.results
              and s_drt.results["combined"].r_pol > 0
              and len(s_drt.results["combined"].peaks) >= 1)
        check("imported DRT curve yields peaks", ok,
              "R_pol=%.2f, %d peak(s)" % (s_drt.results["combined"].r_pol,
                                          len(s_drt.results["combined"].peaks)))
        tot = sum(s_drt.results["combined"].region_r.values())
        err = abs(tot - s_drt.results["combined"].r_pol) / s_drt.results["combined"].r_pol * 100
        check("imported DRT conserves area", err < 0.5, "%.3f %% off" % err)
    except Exception as ex:
        check("imported DRT curve yields peaks", False, str(ex))
        traceback.print_exc()

    # (e) diagnostics module must run clean
    try:
        from drt_studio.core.diagnostics import run_diagnostics, check_environment
        env = check_environment()
        check("environment check runs", env.n_pass > 0,
              "%d ok, %d missing" % (env.n_pass, env.n_fail))
        rep = run_diagnostics(quick=True)
        check("numerical diagnostics all pass", rep.ok,
              "%d passed, %d failed" % (rep.n_pass, rep.n_fail))
        if not rep.ok:
            for c in rep.checks:
                if not c.ok:
                    print("        -> %s: %s" % (c.name, c.detail.splitlines()[0]))
    except Exception as ex:
        check("numerical diagnostics all pass", False, str(ex))
        traceback.print_exc()

    # (f) numpy 1.x compatibility: np.trapezoid does not exist before numpy 2.0
    try:
        from drt_studio.core.compat import trapezoid as _tz
        _x = np.log(np.array([1.0, 2.0, 4.0]))
        _y = np.array([1.0, 1.0, 1.0])
        check("compat.trapezoid works", abs(float(_tz(_y, _x)) - np.log(4.0)) < 1e-9,
              "resolved to np.%s" % _tz.__name__)
    except Exception as ex:
        check("compat.trapezoid works", False, str(ex))

    try:
        import subprocess as _sp
        # re-run the critical numerics in a child process with trapezoid hidden
        code = (
            "import matplotlib; matplotlib.use('Agg')\n"
            "import scipy.optimize, scipy.linalg, scipy.signal, pandas\n"
            "import numpy as np\n"
            "if hasattr(np,'trapezoid'): del np.trapezoid\n"
            "import sys; sys.path.insert(0, %r)\n"
            "from drt_studio.core.synth import make_spectrum\n"
            "from drt_studio.core.model import DRTSettings, Sample\n"
            "from drt_studio.core.pipeline import analyse_sample\n"
            "f,zr,zi = make_spectrum(noise_pct=0.3, seed=1)\n"
            "s = Sample(name='c'); s.f, s.z_re, s.z_im = f, zr, zi\n"
            "analyse_sample(s, DRTSettings(), ('combined',), False, 0.02, lambda m,x: None)\n"
            "r = s.results['combined']\n"
            "assert r.r_pol > 0 and abs(sum(r.region_r.values()) - r.r_pol) < 0.01\n"
            "print('OK', round(r.r_pol, 2))\n"
        ) % os.path.dirname(os.path.abspath(__file__))
        cp = _sp.run([sys.executable, "-c", code], capture_output=True, text=True,
                     timeout=300)
        ok = cp.returncode == 0 and "OK" in cp.stdout
        check("analysis runs without np.trapezoid (numpy 1.x)", ok,
              cp.stdout.strip() or (cp.stderr.strip().splitlines() or [""])[-1][:120])
    except Exception as ex:
        check("analysis runs without np.trapezoid (numpy 1.x)", False, str(ex))

    # (g) every module must still parse as Python 3.8 syntax
    try:
        import ast as _ast
        bad = []
        root = os.path.dirname(os.path.abspath(__file__))
        for dirpath, _dn, fns in os.walk(os.path.join(root, "drt_studio")):
            if "__pycache__" in dirpath:
                continue
            for fn in fns:
                if fn.endswith(".py"):
                    fp = os.path.join(dirpath, fn)
                    try:
                        _ast.parse(open(fp).read(), feature_version=(3, 8))
                    except SyntaxError as se:
                        bad.append("%s:%s" % (fn, se.lineno))
        check("all modules are Python 3.8 compatible", not bad,
              ", ".join(bad) or "syntax OK on 3.8")
    except Exception as ex:
        check("all modules are Python 3.8 compatible", False, str(ex))

    # ------------------------------------------------------------------
    print("\n" + "=" * 78)
    print("RESULT: %d passed, %d failed" % (len(PASS), len(FAIL)))
    if FAIL:
        print("Failed checks:")
        for n in FAIL:
            print("   - " + n)
    print("=" * 78)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
