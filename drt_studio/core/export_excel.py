"""Excel workbook export (openpyxl) - data, DRT curves, peaks, comparison, figures."""
from __future__ import annotations

import os
from typing import Dict, List, Optional

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import MODES, MODE_SHORT, Project, Sample
from .peaks import ZINC_WINDOWS, compare_results, interpret

HDR_FILL = PatternFill("solid", fgColor="1F6390")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13, color="0D2F4C")
SUB_FONT = Font(bold=True, size=11, color="1F6390")
THIN = Side(style="thin", color="B9C6D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, values: List[str], widths: Optional[List[int]] = None):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _rows(ws, start: int, rows: List[list]):
    r = start
    for row in rows:
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            if isinstance(v, float):
                c.number_format = "0.000E+00" if (v and abs(v) < 1e-2) else "0.000"
        r += 1
    return r


def export_workbook(project: Project, path: str, figures: Dict[str, str] = None,
                    include_comparison: bool = True) -> str:
    """Write the complete project to one .xlsx file."""
    figures = figures or {}
    wb = Workbook()

    # ---------------- overview ----------------
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "DRT Studio - project report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = project.name
    ws["A2"].font = SUB_FONT
    ws["A3"] = project.cell_description
    ws["A3"].alignment = WRAP
    ws.column_dimensions["A"].width = 34
    for j in "BCDEFGH":
        ws.column_dimensions[j].width = 16

    r = 5
    _header(ws, r, ["Sample", "Points", "Frequency range", "Modes computed",
                    "R_inf (ohm)", "R_pol (ohm)", "R_total (ohm)", "KK max |res| %"])
    r += 1
    rows = []
    for s in project.samples:
        best = s.best_result()
        rows.append([
            s.name, s.n_points, s.f_range,
            ", ".join(MODE_SHORT[m] for m in MODES if m in s.results) or "-",
            float(best.r_inf) if best else None,
            float(best.r_pol) if best else None,
            float(best.r_inf + best.r_pol) if best else None,
            float(s.kk.max_abs_pct) if s.kk else None,
        ])
    r = _rows(ws, r, rows)
    r += 1
    ws.cell(row=r, column=1, value="Settings").font = SUB_FONT
    r += 1
    ws.cell(row=r, column=1, value=project.settings.describe()).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)

    # ---------------- per sample ----------------
    for s in project.samples:
        safe = _safe(s.name)

        # raw EIS
        if s.has_eis:
            w2 = wb.create_sheet(("EIS_%s" % safe)[:31])
            _header(w2, 1, ["f (Hz)", "Z' (ohm)", "Z'' (ohm)", "-Z'' (ohm)",
                            "|Z| (ohm)", "phase (deg)"],
                    [14, 14, 14, 14, 14, 14])
            ph = np.degrees(np.arctan2(s.z_im, s.z_re))
            data = [[float(a), float(b), float(c), float(-c), float(np.hypot(b, c)), float(d)]
                    for a, b, c, d in zip(s.f, s.z_re, s.z_im, ph)]
            _rows(w2, 2, data)

        # DRT curves
        for m in MODES:
            if m not in s.results:
                continue
            r0 = s.results[m]
            w3 = wb.create_sheet(("DRT_%s_%s" % (MODE_SHORT[m], safe))[:31])
            _header(w3, 1, ["tau (s)", "f (Hz)", "gamma (ohm)"], [16, 16, 16])
            data = [[float(t), float(1 / (2 * np.pi * t)), float(g)]
                    for t, g in zip(r0.tau, r0.gamma)]
            _rows(w3, 2, data)
            ch = LineChart()
            ch.title = "DRT %s - %s" % (MODE_SHORT[m], s.name)
            ch.y_axis.title = "gamma / ohm"
            ch.x_axis.title = "tau / s"
            ch.height, ch.width = 9, 18
            n = len(data)
            ch.add_data(Reference(w3, min_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
            ch.set_categories(Reference(w3, min_col=1, min_row=2, max_row=n + 1))
            w3.add_chart(ch, "E2")

        # peaks + interpretation
        w4 = wb.create_sheet(("Peaks_%s" % safe)[:31])
        rr = 1
        w4.cell(row=rr, column=1, value="Peak table - %s" % s.name).font = TITLE_FONT
        rr += 2
        _header(w4, rr, ["Mode", "Peak", "tau (s)", "f (Hz)", "gamma_max (ohm)",
                         "R_i (ohm)", "% of R_pol", "FWHM (dec)", "Assignment", "Process"],
                [11, 8, 13, 13, 15, 13, 11, 11, 24, 52])
        rr += 1
        rows = []
        for m in MODES:
            if m not in s.results:
                continue
            r0 = s.results[m]
            for p in r0.peaks:
                rows.append([MODE_SHORT[m], p.label, float(p.tau), float(p.freq),
                             float(p.gamma_max), float(p.resistance),
                             float(100 * p.resistance / max(r0.r_pol, 1e-12)),
                             float(p.fwhm_decades), p.assignment, p.process])
        rr = _rows(w4, rr, rows)
        rr += 1

        w4.cell(row=rr, column=1, value="Window resistances").font = SUB_FONT
        rr += 1
        _header(w4, rr, ["Window", "Range (s)", "Meaning"] +
                [MODE_SHORT[m] + " (ohm)" for m in MODES if m in s.results])
        rr += 1
        rows = []
        for w in ZINC_WINDOWS:
            row = [w.key, "%.1e - %.1e" % (w.tau_lo, w.tau_hi), w.short]
            for m in MODES:
                if m in s.results:
                    row.append(float(s.results[m].region_r.get(w.key, 0.0)))
            rows.append(row)
        rr = _rows(w4, rr, rows)
        rr += 1

        w4.cell(row=rr, column=1, value="Automatic interpretation").font = SUB_FONT
        rr += 1
        for m in MODES:
            if m not in s.results:
                continue
            w4.cell(row=rr, column=1, value=MODE_SHORT[m]).font = Font(bold=True)
            rr += 1
            for line in interpret(s.results[m]):
                c = w4.cell(row=rr, column=1, value="- " + line)
                c.alignment = WRAP
                w4.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=10)
                rr += 1
            rr += 1
        for line in compare_results(s.results):
            c = w4.cell(row=rr, column=1, value="* " + line)
            c.alignment = WRAP
            w4.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=10)
            rr += 1

        # KK
        if s.kk is not None and s.kk.res_re_pct.size:
            w5 = wb.create_sheet(("KK_%s" % safe)[:31])
            w5.cell(row=1, column=1, value="Kramers-Kronig test - %s" % s.name).font = TITLE_FONT
            w5.cell(row=2, column=1, value=s.kk.verdict).alignment = WRAP
            w5.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
            _header(w5, 4, ["f (Hz)", "real residual %", "imag residual %"], [16, 18, 18])
            _rows(w5, 5, [[float(a), float(b), float(c)]
                          for a, b, c in zip(s.f, s.kk.res_re_pct, s.kk.res_im_pct)])

        # figures
        figs = figures.get(s.name, [])
        if figs:
            w6 = wb.create_sheet(("Figures_%s" % safe)[:31])
            row = 1
            for fp in figs:
                if os.path.exists(fp):
                    try:
                        img = XLImage(fp)
                        scale = min(1.0, 900.0 / max(img.width, 1))
                        img.width = int(img.width * scale)
                        img.height = int(img.height * scale)
                        w6.add_image(img, "A%d" % row)
                        row += int(img.height / 19) + 3
                    except Exception:
                        pass

    # ---------------- comparison ----------------
    if include_comparison and len(project.samples) > 1:
        wc = wb.create_sheet("Comparison")
        wc.cell(row=1, column=1, value="Sample comparison").font = TITLE_FONT
        r = 3
        for mode in MODES:
            have = [s for s in project.samples if mode in s.results]
            if not have:
                continue
            wc.cell(row=r, column=1,
                    value="Inversion mode: %s" % MODE_SHORT[mode]).font = SUB_FONT
            r += 1
            keys = [w.key for w in ZINC_WINDOWS]
            _header(wc, r, ["Sample", "R_inf (ohm)", "R_pol (ohm)", "R_total (ohm)",
                            "dominant tau (s)", "dominant R (ohm)", "RMS %"] +
                    ["%s (ohm)" % k for k in keys],
                    [22, 13, 13, 13, 15, 15, 10] + [11] * len(keys))
            r += 1
            rows = []
            for s in have:
                res = s.results[mode]
                dom = max(res.peaks, key=lambda p: p.resistance) if res.peaks else None
                rows.append([s.name, float(res.r_inf), float(res.r_pol),
                             float(res.r_inf + res.r_pol),
                             float(dom.tau) if dom else None,
                             float(dom.resistance) if dom else None,
                             float(res.rms_pct)] +
                            [float(res.region_r.get(k, 0.0)) for k in keys])
            r = _rows(wc, r, rows)
            r += 2

        # side-by-side curves
        wd = wb.create_sheet("Comparison_curves")
        col = 1
        for s in project.samples:
            res = s.best_result()
            if res is None:
                continue
            wd.cell(row=1, column=col, value=s.name).font = Font(bold=True)
            wd.cell(row=2, column=col, value="tau (s)").font = HDR_FONT
            wd.cell(row=2, column=col).fill = HDR_FILL
            wd.cell(row=2, column=col + 1, value="gamma (ohm)").font = HDR_FONT
            wd.cell(row=2, column=col + 1).fill = HDR_FILL
            for i, (t, g) in enumerate(zip(res.tau, res.gamma), start=3):
                wd.cell(row=i, column=col, value=float(t))
                wd.cell(row=i, column=col + 1, value=float(g))
            col += 3

        cf = figures.get("__comparison__", [])
        if cf:
            we = wb.create_sheet("Comparison_figures")
            row = 1
            for fp in cf:
                if os.path.exists(fp):
                    try:
                        img = XLImage(fp)
                        scale = min(1.0, 900.0 / max(img.width, 1))
                        img.width = int(img.width * scale)
                        img.height = int(img.height * scale)
                        we.add_image(img, "A%d" % row)
                        row += int(img.height / 19) + 3
                    except Exception:
                        pass

    wb.save(path)
    return path


def _safe(name: str) -> str:
    bad = set(r"[]:*?/\\")
    out = "".join("_" if ch in bad else ch for ch in str(name))
    return out[:18]
