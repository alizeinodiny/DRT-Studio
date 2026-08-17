"""
PDF (ReportLab) and Word (python-docx) report generation.

Two report kinds:
  * single-sample report  - full peak-by-peak analysis of one cell
  * comparison report     - several samples side by side
"""
from __future__ import annotations

import datetime as _dt
import html as _html
import os
from typing import Dict, List, Optional

import numpy as np

from .model import MODES, MODE_SHORT, MODE_LABELS, Project, Sample
from .peaks import ZINC_WINDOWS, compare_results, interpret, ringing_score

# --------------------------------------------------------------------------------------
# shared text blocks
# --------------------------------------------------------------------------------------

TAU = "\u03c4"
GAM = "\u03b3"
OHM = "\u03a9"
LAM = "\u03bb"
APX = "\u2248"

INTRO = (
    "The distribution of relaxation times (DRT) converts an impedance spectrum into a "
    "distribution of characteristic time constants. The impedance is written as a continuum "
    "of parallel RC elements, Z(w) = R_inf + jwL + integral of {} (ln {}) / (1 + jw{}) d ln {}. "
    "Three consequences govern every interpretation in this report: (i) the POSITION of a peak "
    "gives the characteristic frequency f = 1/(2 pi {}) of a process; (ii) the AREA under a peak, "
    "not its height, is its polarisation resistance R = integral {} d ln {}; and (iii) the "
    "inversion is ill-posed and is stabilised by Tikhonov regularisation, so too small a {} "
    "produces spurious oscillations while too large a {} merges genuine peaks."
).format(GAM, TAU, TAU, TAU, TAU, GAM, TAU, LAM, LAM)

MODE_THEORY = (
    "Re-only, Im-only and combined inversions differ because their kernels differ. The imaginary "
    "kernel w{}/(1+w^2{}^2) is a band-pass function peaking sharply at w{} = 1, so an Im-based "
    "inversion is well conditioned and localises processes naturally; it is also blind to any "
    "purely resistive offset. The real kernel 1/(1+w^2{}^2) is a monotone low-pass step in ln {}, "
    "so many different distributions reproduce the same Re(Z) almost equally well: Re-only "
    "inversion is markedly more ill-conditioned and must additionally estimate R_inf from the same "
    "data. Fitting Re and Im simultaneously uses the full information content of the spectrum and "
    "is the recommended reference configuration."
).format(TAU, TAU, TAU, TAU, TAU)

KK_THEORY = (
    "Because Re(Z) and Im(Z) are linked by the Kramers-Kronig relations, a clean, stationary, "
    "causal data set must give essentially the same DRT whichever channel is inverted. A large "
    "disagreement between the three inversions is therefore a data-quality flag, not merely a "
    "numerical curiosity: in aqueous zinc cells the usual cause is drift during the slow "
    "low-frequency sweep (self-discharge, hydrogen evolution, growth of zinc hydroxide sulfate)."
)

CHECKLIST = [
    "Run a linear Kramers-Kronig residual test before any DRT; discard or re-measure any "
    "frequency range with residuals above about 1 %.",
    "Extend the low-frequency limit to 1-10 mHz so that the slowest peak and its tail lie well "
    "inside the {} window, and compute {} one to two decades beyond 1/w_min.".format(TAU, TAU),
    "Use the same {} for every spectrum you intend to compare, chosen objectively by the L-curve "
    "or GCV criterion, and always state it together with the discretisation and RBF width.".format(LAM),
    "Report the combined Re+Im inversion as the primary result; show Re-only and Im-only as a "
    "consistency check, not as three independent measurements.",
    "Repeat the EIS at three temperatures and three states of charge - activation energy and SOC "
    "sensitivity are what turn a peak position into a proven assignment.",
    "Measure a symmetric Zn//Zn cell (and ideally a three-electrode cell) to split the spectrum "
    "into anode and cathode contributions.",
    "Quantify by integrating each peak over ln {}, never by comparing peak heights.".format(TAU),
]


def esc(s) -> str:
    """Escape text that is interpolated into ReportLab paragraph markup."""
    return _html.escape(str(s), quote=False)


def _fmt(v, nd=3):
    if v is None:
        return "-"
    try:
        av = abs(float(v))
    except Exception:
        return str(v)
    if av and (av < 1e-2 or av >= 1e5):
        return ("%." + str(nd) + "e") % v
    return ("%." + str(nd) + "f") % v


def _stamp():
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def _sample_findings(s: Sample) -> List[str]:
    out: List[str] = []
    for m in MODES:
        if m in s.results:
            out.append("[%s] " % MODE_SHORT[m] + " ".join(interpret(s.results[m])[:2]))
    return out


# --------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------

def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping

    try:
        import matplotlib as _mpl
        fd = os.path.join(os.path.dirname(_mpl.__file__), "mpl-data", "fonts", "ttf") + os.sep
        if "DJV" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DJV", fd + "DejaVuSans.ttf"))
            pdfmetrics.registerFont(TTFont("DJV-Bold", fd + "DejaVuSans-Bold.ttf"))
            pdfmetrics.registerFont(TTFont("DJV-Oblique", fd + "DejaVuSans-Oblique.ttf"))
            pdfmetrics.registerFont(TTFont("DJV-BoldOblique", fd + "DejaVuSans-BoldOblique.ttf"))
            addMapping("DJV", 0, 0, "DJV")
            addMapping("DJV", 1, 0, "DJV-Bold")
            addMapping("DJV", 0, 1, "DJV-Oblique")
            addMapping("DJV", 1, 1, "DJV-BoldOblique")
        base, bold, ital = "DJV", "DJV-Bold", "DJV-BoldOblique"
    except Exception:
        base, bold, ital = "Helvetica", "Helvetica-Bold", "Helvetica-BoldOblique"

    ss = getSampleStyleSheet()
    st = {
        "body": ParagraphStyle("body", parent=ss["Normal"], fontName=base, fontSize=9.2,
                               leading=13, alignment=TA_JUSTIFY, spaceAfter=6),
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontName=bold, fontSize=14.5,
                             leading=17.5, textColor=colors.HexColor("#123a5e"),
                             spaceBefore=10, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=bold, fontSize=11,
                             leading=13.5, textColor=colors.HexColor("#1f6390"),
                             spaceBefore=8, spaceAfter=4),
        "h3": ParagraphStyle("h3", parent=ss["Heading3"], fontName=ital, fontSize=9.6,
                             leading=12, textColor=colors.HexColor("#333333"),
                             spaceBefore=6, spaceAfter=3),
        "cap": ParagraphStyle("cap", parent=ss["Normal"], fontName=base, fontSize=8,
                              leading=10.2, alignment=TA_CENTER,
                              textColor=colors.HexColor("#444444"), spaceBefore=3, spaceAfter=10),
        "cell": ParagraphStyle("cell", parent=ss["Normal"], fontName=base, fontSize=7.6,
                               leading=9.5),
        "cellb": ParagraphStyle("cellb", parent=ss["Normal"], fontName=bold, fontSize=7.6,
                                leading=9.5, textColor=colors.white),
        "bullet": ParagraphStyle("bullet", parent=ss["Normal"], fontName=base, fontSize=9.2,
                                 leading=13, leftIndent=12, bulletIndent=3, spaceAfter=3,
                                 alignment=TA_JUSTIFY),
        "title": ParagraphStyle("title", parent=ss["Title"], fontName=bold, fontSize=20,
                                leading=24, textColor=colors.HexColor("#0d2f4c")),
        "sub": ParagraphStyle("sub", parent=ss["Normal"], fontName=base, fontSize=11,
                              leading=14.5, alignment=TA_CENTER,
                              textColor=colors.HexColor("#1f6390")),
    }
    st["_fonts"] = (base, bold)
    return st


def _pdf_table(data, widths, styles, header=True):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    cmds = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9c6d2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3.5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    if header:
        cmds += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f6390")),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                  [colors.white, colors.HexColor("#eef3f8")])]
    t.setStyle(TableStyle(cmds))
    return t


def _pdf_image(path, width):
    from PIL import Image as PILImage
    from reportlab.platypus import Image
    im = PILImage.open(path)
    w, h = im.size
    return Image(path, width=width, height=width * h / w)


def _pdf_doc(path, title, subtitle_left, subtitle_right):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate

    PW, PH = A4

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#1f6390"))
        canvas.setLineWidth(0.7)
        canvas.line(1.6 * cm, PH - 1.5 * cm, PW - 1.6 * cm, PH - 1.5 * cm)
        try:
            canvas.setFont("DJV", 7)
        except Exception:
            canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(1.6 * cm, PH - 1.3 * cm, subtitle_left[:95])
        canvas.drawRightString(PW - 1.6 * cm, PH - 1.3 * cm, subtitle_right[:60])
        canvas.line(1.6 * cm, 1.3 * cm, PW - 1.6 * cm, 1.3 * cm)
        canvas.drawCentredString(PW / 2, 0.92 * cm, "Page %d" % doc.page)
        canvas.restoreState()

    doc = BaseDocTemplate(path, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                          topMargin=1.9 * cm, bottomMargin=1.55 * cm, title=title,
                          author="DRT Studio")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="n")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame], onPage=deco)])
    return doc


def sample_report_pdf(sample: Sample, project: Project, path: str,
                      figures: Dict[str, str], windows=None) -> str:
    """Full analysis report for one sample."""
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, Spacer, KeepTogether

    windows = windows or ZINC_WINDOWS
    st = _pdf_styles()
    doc = _pdf_doc(path, "DRT report - %s" % sample.name,
                   "DRT Studio - %s" % project.name, sample.name)
    P = lambda t, s="body": Paragraph(t, st[s])
    F = []

    # cover
    F.append(Spacer(1, 0.7 * cm))
    F.append(P("Distribution of Relaxation Times<br/>Analysis Report", "title"))
    F.append(Spacer(1, 3 * 2))
    F.append(P(esc(sample.name), "sub"))
    F.append(Spacer(1, 4))
    F.append(P(esc(project.cell_description), "cap"))
    F.append(P("Generated %s%s" % (_stamp(),
               (" by " + project.operator) if project.operator else ""), "cap"))

    # summary table
    rows = [[Paragraph(x, st["cellb"]) for x in
             ["Quantity"] + [MODE_SHORT[m] for m in MODES if m in sample.results]]]
    have = [m for m in MODES if m in sample.results]

    def add(label, fn):
        rows.append([Paragraph(label, st["cell"])] +
                    [Paragraph(fn(sample.results[m]), st["cell"]) for m in have])

    add("Series resistance R_inf (%s)" % OHM, lambda r: _fmt(r.r_inf, 2))
    add("Polarisation resistance R_pol (%s)" % OHM, lambda r: _fmt(r.r_pol, 2))
    add("Total resistance (%s)" % OHM, lambda r: _fmt(r.r_inf + r.r_pol, 2))
    add("Peaks resolved", lambda r: str(len(r.peaks)))
    add("Dominant %s (s)" % TAU,
        lambda r: _fmt(max(r.peaks, key=lambda p: p.resistance).tau, 3) if r.peaks else "-")
    add("Its resistance (%s)" % OHM,
        lambda r: _fmt(max(r.peaks, key=lambda p: p.resistance).resistance, 2) if r.peaks else "-")
    add("Regularisation %s" % LAM, lambda r: _fmt(r.lambda_used))
    add("RMS residual (%)", lambda r: _fmt(r.rms_pct, 2))
    if have:
        w0 = 6.2 * cm
        wr = (doc.width - w0) / len(have)
        F.append(_pdf_table(rows, [w0] + [wr] * len(have), st))
        F.append(P("<b>Table 1.</b> Summary of the inversions computed for this sample.", "cap"))

    # 1 method
    F.append(P("1. &nbsp;Method and settings", "h1"))
    F.append(P(INTRO))
    F.append(P(MODE_THEORY))
    F.append(P("<b>Settings used:</b> " + esc(project.settings.describe())))
    if sample.source_files:
        F.append(P("<b>Source files:</b> " + esc("; ".join(os.path.basename(x)
                                                            for x in sample.source_files))))
    if sample.has_eis:
        F.append(P("<b>Data:</b> %d frequencies, %s." % (sample.n_points, sample.f_range)))
    if sample.notes:
        F.append(P("<b>Notes:</b> " + esc(sample.notes)))

    # 2 data quality
    F.append(P("2. &nbsp;Data quality (Kramers-Kronig)", "h1"))
    F.append(P(KK_THEORY))
    if sample.kk is not None:
        F.append(P("<b>Result:</b> %s" % esc(sample.kk.verdict)))
        F.append(P("Maximum |residual| %.2f %%, mean %.2f %%, %d RC elements."
                   % (sample.kk.max_abs_pct, sample.kk.mean_abs_pct, sample.kk.n_rc)))
        if "kk" in figures:
            F.append(_pdf_image(figures["kk"], doc.width * 0.8))
            F.append(P("<b>Figure.</b> Linear Kramers-Kronig residuals. The shaded band is "
                       "the +/-1 % acceptance corridor.", "cap"))
    else:
        F.append(P("<i>The Kramers-Kronig test was not run for this sample.</i>"))

    F.append(PageBreak())

    # 3 figures
    F.append(P("3. &nbsp;DRT results", "h1"))
    F.append(P("Each inversion is plotted in its own panel on a shared time-constant axis so "
               "that no curve is crushed by another. The coloured strip carries the assignment "
               "windows and applies to all panels; reading down a vertical line compares what "
               "the inversions say about the same process."))
    if "panels" in figures:
        F.append(_pdf_image(figures["panels"], doc.width))
        F.append(P("<b>Figure 1.</b> DRT of %s - one panel per inversion mode, full scale."
                   % sample.name, "cap"))
    if "panels_zoom" in figures:
        F.append(PageBreak())
        F.append(_pdf_image(figures["panels_zoom"], doc.width))
        F.append(P("<b>Figure 2.</b> The same panels on a common magnified scale so the fast "
                   "processes become legible. Peaks running off the top are labelled with an "
                   "arrow and their true height.", "cap"))
    for key, cap in (("nyquist", "Nyquist plot: points are the measured data, lines the "
                                 "impedance reconstructed from the DRT solution."),
                     ("residuals", "Relative residuals of the DRT fit."),
                     ("bars", "Polarisation resistance carried by each assignment window.")):
        if key in figures:
            F.append(_pdf_image(figures[key], doc.width * (0.82 if key != "bars" else 0.92)))
            F.append(P("<b>Figure.</b> " + cap, "cap"))

    F.append(PageBreak())

    # 4 peak table
    F.append(P("4. &nbsp;Peak table", "h1"))
    hdr = ["Mode", "Peak", "%s (s)" % TAU, "f (Hz)", "%s max (%s)" % (GAM, OHM),
           "R_i (%s)" % OHM, "% R_pol", "FWHM", "Assignment"]
    rows = [[Paragraph(h, st["cellb"]) for h in hdr]]
    for m in have:
        r = sample.results[m]
        for p in r.peaks:
            rows.append([Paragraph(esc(x), st["cell"]) for x in [
                MODE_SHORT[m], p.label, _fmt(p.tau, 3), _fmt(p.freq, 3),
                _fmt(p.gamma_max, 2), _fmt(p.resistance, 2),
                "%.0f" % (100 * p.resistance / max(r.r_pol, 1e-12)),
                "%.2f" % p.fwhm_decades, p.assignment]])
    if len(rows) > 1:
        wds = [1.5, 1.1, 1.9, 1.9, 1.9, 1.7, 1.2, 1.1, 4.4]
        sc = doc.width / (sum(wds) * cm)
        F.append(_pdf_table(rows, [w * cm * sc for w in wds], st))
        F.append(P("<b>Table 2.</b> All peaks detected in every inversion, with the physical "
                   "window each falls into.", "cap"))
    else:
        F.append(P("<i>No peaks were detected. Lower the peak-prominence threshold or reduce "
                   "the regularisation.</i>"))

    # window table
    F.append(P("4.1 &nbsp;Resistance per assignment window", "h2"))
    hdr = ["Window", "%s range (s)" % TAU, "Meaning"] + ["%s (%s)" % (MODE_SHORT[m], OHM)
                                                         for m in have]
    rows = [[Paragraph(h, st["cellb"]) for h in hdr]]
    for w in windows:
        row = [w.key, "%.0e - %.0e" % (w.tau_lo, w.tau_hi), w.short]
        for m in have:
            row.append(_fmt(sample.results[m].region_r.get(w.key, 0.0), 2))
        rows.append([Paragraph(esc(x), st["cell"]) for x in row])
    wds = [1.3, 2.6, 5.2] + [1.9] * len(have)
    sc = doc.width / (sum(wds) * cm)
    F.append(_pdf_table(rows, [w * cm * sc for w in wds], st))
    F.append(P("<b>Table 3.</b> Integrated resistance per window, R_i = integral of %s d ln %s."
               % (GAM, TAU), "cap"))

    F.append(PageBreak())

    # 5 assignment
    F.append(P("5. &nbsp;Peak-by-peak physical assignment", "h1"))
    F.append(P("The assignments below follow the time-constant windows for an aqueous zinc-ion "
               "cell with a galvanized zinc anode, a hard-carbon cathode and a mildly acidic "
               "ZnSO4 electrolyte. Each is a testable hypothesis, not a fact read off a single "
               "spectrum; the discriminating experiment is given in each case."))
    for w in windows:
        present = []
        for m in have:
            for p in sample.results[m].peaks:
                if p.label.startswith(w.key):
                    present.append((m, p))
        head = "%s &nbsp;- &nbsp;%s &nbsp;(%s = %.0e - %.0e s)" % (w.key, w.short, TAU,
                                                                   w.tau_lo, w.tau_hi)
        F.append(P(head, "h2"))
        F.append(P("<b>Process.</b> " + esc(w.process)))
        F.append(P(esc(w.detail)))
        if present:
            bits = ["%s: %s at %s = %s s carrying %s %s" %
                    (MODE_SHORT[m], p.label, TAU, _fmt(p.tau, 3), _fmt(p.resistance, 2), OHM)
                    for m, p in present]
            F.append(P("<b>Observed here.</b> " + "; ".join(bits) + "."))
        else:
            tot = sum(sample.results[m].region_r.get(w.key, 0.0) for m in have) / max(len(have), 1)
            if tot > 0.05:
                F.append(P("<b>Observed here.</b> No separate maximum, but this window still "
                           "carries about %s %s on average - it appears as part of a distributed "
                           "plateau rather than a discrete peak." % (_fmt(tot, 2), OHM)))
            else:
                F.append(P("<b>Observed here.</b> No measurable contribution in this window."))

    # 6 findings
    F.append(PageBreak())
    F.append(P("6. &nbsp;Automatic findings and caveats", "h1"))
    for m in have:
        F.append(P("6.%d &nbsp;%s" % (have.index(m) + 1, MODE_LABELS[m]), "h2"))
        for line in interpret(sample.results[m]):
            F.append(Paragraph(esc(line), st["bullet"], bulletText="\u2022"))
    cmp_lines = compare_results(sample.results)
    if cmp_lines:
        F.append(P("6.%d &nbsp;Cross-check between inversions" % (len(have) + 1), "h2"))
        for line in cmp_lines:
            F.append(Paragraph(esc(line), st["bullet"], bulletText="\u2022"))

    F.append(P("7. &nbsp;Measurement checklist", "h1"))
    for i, line in enumerate(CHECKLIST, 1):
        F.append(Paragraph(line, st["bullet"], bulletText="%d." % i))

    doc.build(F)
    return path


def comparison_report_pdf(project: Project, samples: List[Sample], path: str,
                          figures: Dict[str, str], mode: str = "combined",
                          windows=None) -> str:
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, Spacer

    windows = windows or ZINC_WINDOWS
    st = _pdf_styles()
    doc = _pdf_doc(path, "DRT comparison report", "DRT Studio - %s" % project.name,
                   "Comparison of %d samples" % len(samples))
    P = lambda t, s="body": Paragraph(t, st[s])
    F = []

    F.append(Spacer(1, 0.7 * cm))
    F.append(P("DRT Comparison Report", "title"))
    F.append(Spacer(1, 6))
    F.append(P("%d samples compared using the %s inversion"
               % (len(samples), MODE_SHORT.get(mode, mode)), "sub"))
    F.append(Spacer(1, 4))
    F.append(P(esc(project.cell_description), "cap"))
    F.append(P("Generated %s%s" % (_stamp(),
               (" by " + project.operator) if project.operator else ""), "cap"))

    # master table
    keys = [w.key for w in windows]
    hdr = ["Sample", "R_inf", "R_pol", "R_tot", "dom. %s (s)" % TAU, "dom. R", "RMS %"] + keys
    rows = [[Paragraph(h, st["cellb"]) for h in hdr]]
    for s in samples:
        r = s.results.get(mode)
        if r is None:
            continue
        dom = max(r.peaks, key=lambda p: p.resistance) if r.peaks else None
        row = [s.name, _fmt(r.r_inf, 2), _fmt(r.r_pol, 1), _fmt(r.r_inf + r.r_pol, 1),
               _fmt(dom.tau, 3) if dom else "-", _fmt(dom.resistance, 1) if dom else "-",
               _fmt(r.rms_pct, 2)] + [_fmt(r.region_r.get(k, 0.0), 1) for k in keys]
        rows.append([Paragraph(esc(x), st["cell"]) for x in row])
    wds = [3.0, 1.3, 1.3, 1.3, 1.7, 1.3, 1.1] + [1.0] * len(keys)
    sc = doc.width / (sum(wds) * cm)
    F.append(_pdf_table(rows, [w * cm * sc for w in wds], st))
    F.append(P("<b>Table 1.</b> Master comparison. All resistances in %s; window columns are "
               "the integrated resistance R_i of each assignment window." % OHM, "cap"))

    # ranking / findings
    F.append(P("1. &nbsp;What the comparison shows", "h1"))
    valid = [s for s in samples if mode in s.results]
    if valid:
        by_tot = sorted(valid, key=lambda s: s.results[mode].r_inf + s.results[mode].r_pol)
        best, worst = by_tot[0], by_tot[-1]
        rb = best.results[mode].r_inf + best.results[mode].r_pol
        rw = worst.results[mode].r_inf + worst.results[mode].r_pol
        F.append(P("<b>Lowest total resistance:</b> %s (%s %s). <b>Highest:</b> %s (%s %s), "
                   "a factor %.2f higher."
                   % (best.name, _fmt(rb, 1), OHM, worst.name, _fmt(rw, 1), OHM,
                      rw / max(rb, 1e-12))))
        # which window drives the difference
        diffs = []
        for w in windows:
            a = best.results[mode].region_r.get(w.key, 0.0)
            b = worst.results[mode].region_r.get(w.key, 0.0)
            diffs.append((abs(b - a), w, a, b))
        diffs.sort(reverse=True, key=lambda x: x[0])
        d, w, a, b = diffs[0]
        F.append(P("<b>Dominant difference:</b> window %s (%s) accounts for %s %s of the gap "
                   "(%s %s in %s versus %s %s in %s). %s"
                   % (w.key, w.short, _fmt(d, 1), OHM, _fmt(a, 1), OHM, best.name,
                      _fmt(b, 1), OHM, worst.name, w.process)))
        # spread per window
        F.append(P("1.1 &nbsp;Spread per window", "h2"))
        for w in windows:
            vals = [s.results[mode].region_r.get(w.key, 0.0) for s in valid]
            if max(vals) < 0.05:
                continue
            F.append(Paragraph(
                "<b>%s (%s):</b> %s to %s %s across the set (mean %s, spread factor %.1f). %s"
                % (w.key, w.short, _fmt(min(vals), 2), _fmt(max(vals), 2), OHM,
                   _fmt(float(np.mean(vals)), 2), max(vals) / max(min(vals), 1e-9), w.short),
                st["bullet"], bulletText="\u2022"))

    # figures
    F.append(PageBreak())
    F.append(P("2. &nbsp;Comparison figures", "h1"))
    order = [("stacked", "Each sample in its own panel on a shared time-constant axis - the "
                         "clearest way to see which process differs."),
             ("overlay", "All samples overlaid on one axes."),
             ("bars", "Polarisation resistance per assignment window."),
             ("heatmap", "Process fingerprint: the share of R_pol in each window."),
             ("summary", "Resistance budget and dominant process."),
             ("nyquist", "Nyquist plots of the raw data.")]
    for key, cap in order:
        if key in figures:
            F.append(_pdf_image(figures[key], doc.width * (1.0 if key == "stacked" else 0.86)))
            F.append(P("<b>Figure.</b> " + cap, "cap"))

    # per-sample digest
    F.append(PageBreak())
    F.append(P("3. &nbsp;Per-sample digest", "h1"))
    for s in samples:
        r = s.results.get(mode)
        if r is None:
            continue
        F.append(P(esc(s.name), "h2"))
        if s.notes:
            F.append(P("<i>%s</i>" % esc(s.notes)))
        for line in interpret(r):
            F.append(Paragraph(esc(line), st["bullet"], bulletText="\u2022"))

    F.append(P("4. &nbsp;Measurement checklist", "h1"))
    for i, line in enumerate(CHECKLIST, 1):
        F.append(Paragraph(line, st["bullet"], bulletText="%d." % i))

    doc.build(F)
    return path


# --------------------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------------------

def _docx_base(title: str, subtitle: str, project: Project):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    d = Document()
    for name, size in (("Normal", 10),):
        style = d.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)

    h = d.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(subtitle)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(12)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x63, 0x90)
    p2 = d.add_paragraph(project.cell_description)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.runs[0].font.size = Pt(9)
    p3 = d.add_paragraph("Generated %s%s" % (_stamp(),
                         (" by " + project.operator) if project.operator else ""))
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.runs[0].font.size = Pt(8)
    return d


def _docx_table(doc, header: List[str], rows: List[List[str]]):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        for par in cell.paragraphs:
            for run in par.runs:
                run.font.bold = True
                run.font.size = __import__("docx").shared.Pt(8)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = __import__("docx").shared.Pt(8)
    return t


def sample_report_docx(sample: Sample, project: Project, path: str,
                       figures: Dict[str, str], windows=None) -> str:
    from docx.shared import Inches

    windows = windows or ZINC_WINDOWS
    d = _docx_base("Distribution of Relaxation Times - Analysis Report",
                   sample.name, project)
    have = [m for m in MODES if m in sample.results]

    d.add_heading("Summary", level=1)
    rows = []
    labels = [
        ("Series resistance R_inf (ohm)", lambda r: _fmt(r.r_inf, 2)),
        ("Polarisation resistance R_pol (ohm)", lambda r: _fmt(r.r_pol, 2)),
        ("Total resistance (ohm)", lambda r: _fmt(r.r_inf + r.r_pol, 2)),
        ("Peaks resolved", lambda r: str(len(r.peaks))),
        ("Dominant tau (s)",
         lambda r: _fmt(max(r.peaks, key=lambda p: p.resistance).tau, 3) if r.peaks else "-"),
        ("Regularisation lambda", lambda r: _fmt(r.lambda_used)),
        ("RMS residual (%)", lambda r: _fmt(r.rms_pct, 2)),
    ]
    for lab, fn in labels:
        rows.append([lab] + [fn(sample.results[m]) for m in have])
    _docx_table(d, ["Quantity"] + [MODE_SHORT[m] for m in have], rows)

    d.add_heading("1. Method and settings", level=1)
    d.add_paragraph(INTRO)
    d.add_paragraph(MODE_THEORY)
    d.add_paragraph("Settings: " + project.settings.describe())
    if sample.source_files:
        d.add_paragraph("Source files: " + "; ".join(os.path.basename(x)
                                                     for x in sample.source_files))
    if sample.has_eis:
        d.add_paragraph("Data: %d frequencies, %s." % (sample.n_points, sample.f_range))
    if sample.notes:
        d.add_paragraph("Notes: " + sample.notes)

    d.add_heading("2. Data quality (Kramers-Kronig)", level=1)
    d.add_paragraph(KK_THEORY)
    if sample.kk is not None:
        d.add_paragraph("Result: " + sample.kk.verdict)
        if "kk" in figures and os.path.exists(figures["kk"]):
            d.add_picture(figures["kk"], width=Inches(5.6))
    else:
        d.add_paragraph("The Kramers-Kronig test was not run for this sample.")

    d.add_heading("3. DRT results", level=1)
    d.add_paragraph("Each inversion is plotted in its own panel on a shared time-constant axis. "
                    "The coloured strip carries the assignment windows and applies to all panels.")
    for key, cap in (("panels", "Figure 1. One panel per inversion mode, full scale."),
                     ("panels_zoom", "Figure 2. The same panels, magnified scale."),
                     ("nyquist", "Figure. Nyquist plot with the DRT reconstruction."),
                     ("residuals", "Figure. Relative residuals of the DRT fit."),
                     ("bars", "Figure. Resistance per assignment window.")):
        if key in figures and os.path.exists(figures[key]):
            d.add_picture(figures[key], width=Inches(6.3 if "panels" in key else 5.4))
            cp = d.add_paragraph(cap)
            cp.runs[0].font.size = __import__("docx").shared.Pt(8)

    d.add_heading("4. Peak table", level=1)
    rows = []
    for m in have:
        r = sample.results[m]
        for p in r.peaks:
            rows.append([MODE_SHORT[m], p.label, _fmt(p.tau, 3), _fmt(p.freq, 3),
                         _fmt(p.gamma_max, 2), _fmt(p.resistance, 2),
                         "%.0f" % (100 * p.resistance / max(r.r_pol, 1e-12)),
                         p.assignment])
    if rows:
        _docx_table(d, ["Mode", "Peak", "tau (s)", "f (Hz)", "gamma max", "R_i (ohm)",
                        "% R_pol", "Assignment"], rows)

    d.add_heading("4.1 Resistance per assignment window", level=2)
    rows = []
    for w in windows:
        rows.append([w.key, "%.0e - %.0e" % (w.tau_lo, w.tau_hi), w.short] +
                    [_fmt(sample.results[m].region_r.get(w.key, 0.0), 2) for m in have])
    _docx_table(d, ["Window", "tau range (s)", "Meaning"] +
                [MODE_SHORT[m] + " (ohm)" for m in have], rows)

    d.add_heading("5. Peak-by-peak physical assignment", level=1)
    for w in windows:
        d.add_heading("%s - %s (tau = %.0e - %.0e s)" % (w.key, w.short, w.tau_lo, w.tau_hi),
                      level=2)
        d.add_paragraph("Process: " + w.process)
        d.add_paragraph(w.detail)
        present = []
        for m in have:
            for p in sample.results[m].peaks:
                if p.label.startswith(w.key):
                    present.append("%s: %s at tau = %s s carrying %s ohm"
                                   % (MODE_SHORT[m], p.label, _fmt(p.tau, 3),
                                      _fmt(p.resistance, 2)))
        d.add_paragraph("Observed here: " + ("; ".join(present) if present
                                             else "no discrete peak in this window."))

    d.add_heading("6. Automatic findings", level=1)
    for m in have:
        d.add_heading(MODE_LABELS[m], level=2)
        for line in interpret(sample.results[m]):
            d.add_paragraph(line, style="List Bullet")
    cl = compare_results(sample.results)
    if cl:
        d.add_heading("Cross-check between inversions", level=2)
        for line in cl:
            d.add_paragraph(line, style="List Bullet")

    d.add_heading("7. Measurement checklist", level=1)
    for line in CHECKLIST:
        d.add_paragraph(line, style="List Number")

    d.save(path)
    return path


def comparison_report_docx(project: Project, samples: List[Sample], path: str,
                           figures: Dict[str, str], mode: str = "combined",
                           windows=None) -> str:
    from docx.shared import Inches

    windows = windows or ZINC_WINDOWS
    d = _docx_base("DRT Comparison Report",
                   "%d samples compared using the %s inversion"
                   % (len(samples), MODE_SHORT.get(mode, mode)), project)

    keys = [w.key for w in windows]
    rows = []
    for s in samples:
        r = s.results.get(mode)
        if r is None:
            continue
        dom = max(r.peaks, key=lambda p: p.resistance) if r.peaks else None
        rows.append([s.name, _fmt(r.r_inf, 2), _fmt(r.r_pol, 1),
                     _fmt(r.r_inf + r.r_pol, 1),
                     _fmt(dom.tau, 3) if dom else "-",
                     _fmt(r.rms_pct, 2)] + [_fmt(r.region_r.get(k, 0.0), 1) for k in keys])
    _docx_table(d, ["Sample", "R_inf", "R_pol", "R_tot", "dom. tau (s)", "RMS %"] + keys, rows)

    d.add_heading("1. What the comparison shows", level=1)
    valid = [s for s in samples if mode in s.results]
    if valid:
        by_tot = sorted(valid, key=lambda s: s.results[mode].r_inf + s.results[mode].r_pol)
        best, worst = by_tot[0], by_tot[-1]
        rb = best.results[mode].r_inf + best.results[mode].r_pol
        rw = worst.results[mode].r_inf + worst.results[mode].r_pol
        d.add_paragraph("Lowest total resistance: %s (%s ohm). Highest: %s (%s ohm), a factor "
                        "%.2f higher." % (best.name, _fmt(rb, 1), worst.name, _fmt(rw, 1),
                                          rw / max(rb, 1e-12)))
        for w in windows:
            vals = [s.results[mode].region_r.get(w.key, 0.0) for s in valid]
            if max(vals) < 0.05:
                continue
            d.add_paragraph("%s (%s): %s to %s ohm across the set. %s"
                            % (w.key, w.short, _fmt(min(vals), 2), _fmt(max(vals), 2),
                               w.process), style="List Bullet")

    d.add_heading("2. Comparison figures", level=1)
    for key, cap in (("stacked", "Each sample in its own panel on a shared axis."),
                     ("overlay", "All samples overlaid."),
                     ("bars", "Resistance per assignment window."),
                     ("heatmap", "Process fingerprint (share of R_pol)."),
                     ("summary", "Resistance budget and dominant process."),
                     ("nyquist", "Nyquist plots of the raw data.")):
        if key in figures and os.path.exists(figures[key]):
            d.add_picture(figures[key], width=Inches(6.3 if key == "stacked" else 5.6))
            cp = d.add_paragraph(cap)
            cp.runs[0].font.size = __import__("docx").shared.Pt(8)

    d.add_heading("3. Per-sample digest", level=1)
    for s in samples:
        r = s.results.get(mode)
        if r is None:
            continue
        d.add_heading(s.name, level=2)
        if s.notes:
            d.add_paragraph(s.notes)
        for line in interpret(r):
            d.add_paragraph(line, style="List Bullet")

    d.add_heading("4. Measurement checklist", level=1)
    for line in CHECKLIST:
        d.add_paragraph(line, style="List Number")

    d.save(path)
    return path
